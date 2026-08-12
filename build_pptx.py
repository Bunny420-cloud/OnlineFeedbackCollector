import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Define Color Palette
COLOR_BG_DARK = RGBColor(15, 23, 42)        # #0f172a (Navy Dark)
COLOR_CARD_BG = RGBColor(22, 30, 49)        # #161e31 (Glass Card Dark)
COLOR_CARD_BORDER = RGBColor(40, 50, 75)    # Subtle border
COLOR_PRIMARY = RGBColor(99, 102, 241)      # #6366f1 (Indigo Accent)
COLOR_CYAN = RGBColor(56, 189, 248)         # #38bdf8 (Cyan Accent)
COLOR_GOLD = RGBColor(251, 191, 36)         # #fbbf24 (Gold Accent)
COLOR_TEXT_WHITE = RGBColor(255, 255, 255)  # White Title Text
COLOR_TEXT_LIGHT = RGBColor(203, 213, 225)  # #cbd5e1 Light Text
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8 Muted Text

def apply_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG_DARK

def add_header(slide, section_tag, slide_title):
    # Tag Pill
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(6), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = section_tag.upper()
    p_tag.font.name = 'Georgia'
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_CYAN

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = slide_title
    p_title.font.name = 'Georgia'
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE

def add_card(slide, left, top, width, height, title, items, badge_text=None):
    # Card Background Shape
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD_BG
    shape.line.color.rgb = COLOR_CARD_BORDER
    shape.line.width = Pt(1)

    # Card Content Textbox
    tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = 'Georgia'
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_CYAN
        p0.space_after = Pt(10)

    for item in items:
        p = tf.add_paragraph()
        p.text = f"• {item}" if not title else item
        p.font.name = 'Arial'
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        p.space_after = Pt(6)

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_background(slide1)

    # Title Hero Box
    tb_hero = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(2.2))
    tf_hero = tb_hero.text_frame
    tf_hero.word_wrap = True

    p0 = tf_hero.paragraphs[0]
    p0.text = "A SUMMER INTERNSHIP PROJECT PRESENTATION"
    p0.font.name = 'Arial'
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN
    p0.space_after = Pt(10)

    p1 = tf_hero.add_paragraph()
    p1.text = "Online Feedback Collector with Admin Dashboard"
    p1.font.name = 'Georgia'
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE
    p1.space_after = Pt(10)

    p2 = tf_hero.add_paragraph()
    p2.text = "Python Developer Internship | Full-Stack Web Application & Visual Analytics System"
    p2.font.name = 'Arial'
    p2.font.size = Pt(15)
    p2.font.color.rgb = COLOR_TEXT_LIGHT

    # Metadata Cards (2 Columns)
    add_card(slide1, 1.0, 3.8, 5.4, 3.0, "Presented By", [
        "Student Name: Munukuntla Bunny Narayana Naidu",
        "Roll / Enrollment No.: 23CS002707",
        "Program: B.Tech (Computer Science & Engineering)",
        "Institution: Sir Padampat Singhania University (SPSU), Udaipur",
        "Department: Department of Computing and Informatics"
    ])

    add_card(slide1, 6.9, 3.8, 5.4, 3.0, "Internship Mentorship Details", [
        "Host Organization: The Skybrisk (www.theskybrisk.com)",
        "Project Guide: Govind Bhokare (The Skybrisk)",
        "Internship Role: Python Developer Intern",
        "Duration: 2 Months (May 20, 2026 – July 20, 2026)",
        "Certificate Ref ID: EMP20250320-03229"
    ])

    # =========================================================================
    # SLIDE 2: INTRODUCTION OF COMPANY / ORGANIZATION
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_background(slide2)
    add_header(slide2, "Section 1: Company Profile", "Introduction of The Skybrisk")

    add_card(slide2, 0.8, 1.6, 5.6, 5.2, "Company Background & Overview", [
        "Organization Name: The Skybrisk",
        "Motto: 'Where Technology Meets Creativity'",
        "Official Website: https://www.theskybrisk.com",
        "Primary Focus: Next-generation technology consulting, custom software engineering, and digital transformation.",
        "Company Culture: Emphasizes agile software methodologies, practical project-based execution, and high-performance cloud solutions.",
        "Internship Engagement: Provides rigorous Python developer training, industry code reviews, and enterprise solution development."
    ])

    add_card(slide2, 6.8, 1.6, 5.6, 5.2, "Major Products & Industry Domain", [
        "Custom Web & Enterprise Applications: Building scalable web microservices and enterprise software portals.",
        "Cloud & API Solutions: Architecting RESTful microservices, database migration pipelines, and serverless backends.",
        "Artificial Intelligence & Analytics: Developing data analytics dashboards, machine learning solutions, and business intelligence tools.",
        "Industry Domain: Software Engineering, Web Platform Architecture, Data Management, and Information Systems."
    ])

    # =========================================================================
    # SLIDE 3: INTRODUCTION OF THE PROJECT
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_background(slide3)
    add_header(slide3, "Section 2: Project Introduction", "Project Overview & Problem Statement")

    add_card(slide3, 0.8, 1.6, 3.7, 5.2, "Project Title & Context", [
        "Title: Online Feedback Collector with Admin Dashboard",
        "Domain: Full-Stack Web Development & Data Analytics",
        "Core Tech: Python 3, Flask, SQLite, Bootstrap 5, Chart.js",
        "Target Audience: End-users submitting feedback & System Administrators monitoring operational quality."
    ])

    add_card(slide3, 4.8, 1.6, 3.7, 5.2, "Problem Statement", [
        "Traditional Surveys cause annoying full-page reloads, leading to form abandonment.",
        "Fragmented Feedback Data makes manual analysis slow and prone to errors.",
        "Lack of Real-Time Analytics prevents decision-makers from spotting issues fast.",
        "Security Gaps in plain static forms leave feedback data unprotected."
    ])

    add_card(slide3, 8.8, 1.6, 3.7, 5.2, "Purpose & Key Objectives", [
        "Build a lightweight, zero-reload 5-star feedback collection web interface.",
        "Engineered an SQLite serverless database schema with automatic startup setup.",
        "Develop an interactive Admin Dashboard featuring dynamic Chart.js analytics.",
        "Provide one-click CSV export and a clean public REST API endpoint."
    ])

    # =========================================================================
    # SLIDE 4: BRIEF IDEA ABOUT THE PROJECT
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_background(slide4)
    add_header(slide4, "Section 3: Brief Idea", "Proposed Work, Methodology & Expected Outcomes")

    add_card(slide4, 0.8, 1.6, 3.7, 5.2, "Overview of Proposed Work", [
        "Decoupled Client-Server Model: HTML5/CSS3/JS frontend communicating with Python Flask WSGI server.",
        "Interactive Star Widget: 5-star rating widget with hover states and dynamic preview labels.",
        "Session-Based Security: Secure admin login portal with hashed password authentication (`Werkzeug`)."
    ])

    add_card(slide4, 4.8, 1.6, 3.7, 5.2, "Methodology & Approach", [
        "Client-Side Validation: Instant JavaScript checks on name, email format, rating selection, and comments length.",
        "AJAX Async Submission: Native `fetch` API POST requests to `/submit-feedback` without page refresh.",
        "Data Persistence & Analytics: SQL queries over SQLite (`database.db`) feeding Chart.js graphs."
    ])

    add_card(slide4, 8.8, 1.6, 3.7, 5.2, "Expected Outcomes", [
        "Sub-15ms REST API latency across all CRUD endpoints.",
        "Instant visual rendering of Rating Distribution & Submission Trend charts.",
        "Zero-lag client-side live table search.",
        "One-click CSV report downloading."
    ])

    # =========================================================================
    # SLIDE 5: NOVELTY OF THE PROJECT
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_background(slide5)
    add_header(slide5, "Section 4: Novelty & Innovation", "Innovative Aspects & Key Value Addition")

    add_card(slide5, 0.8, 1.6, 3.7, 5.2, "Innovative Aspects", [
        "Interactive Glassmorphism UI: Custom dark-mode design tokens with glowing focus rings and smooth transitions.",
        "Zero-Reload AJAX Lifecycle: Completely eliminates page refreshes during submission, improving user completion rates.",
        "Star-Rating Badge Sync: Real-time visual feedback syncing radio button inputs with star badges."
    ])

    add_card(slide5, 4.8, 1.6, 3.7, 5.2, "Differentiation from Existing Tools", [
        "Lightweight & Fast: Built using Flask micro-framework instead of heavy, resource-intensive SPA frameworks.",
        "Zero-Config Database: SQLite persistent storage removes complex external database administration.",
        "Integrated Analytics: Combines form collection, live charts, and admin table filtering into a single unified web solution."
    ])

    add_card(slide5, 8.8, 1.6, 3.7, 5.2, "Key Value Addition", [
        "Executive Decision Support: Live KPI summary cards (Total Feedback, Avg Rating, 5-Star & 1-Star Counts).",
        "Interoperable REST API: `GET /api/feedback` allows easy integration with third-party mobile or web systems.",
        "RFC 4180 CSV Export: Structured data downloads for offline analysis in Excel or Google Sheets."
    ])

    # =========================================================================
    # SLIDE 6: PROJECT DESCRIPTION & TECH STACK
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_background(slide6)
    add_header(slide6, "Section 5: Project Description", "Detailed Architecture, Tech Stack & Workflow")

    add_card(slide6, 0.8, 1.6, 5.6, 5.2, "Technologies & Tools Used", [
        "Backend Framework: Python 3.12 + Flask 3.0.2 micro-framework",
        "Database Engine: SQLite3 (`database.db`) with auto DDL initialization",
        "Security & Auth: Werkzeug password hashing + Flask Session management",
        "Frontend Layout: HTML5, Jinja2 Templating, Bootstrap 5.3.2 grid",
        "Custom Styling: CSS3 Glassmorphism, CSS variables, FontAwesome 6.5",
        "Data Visualizations: Chart.js 4.4.1 (Bar & Line chart canvases)",
        "Async Communication: JavaScript ES6+ native `fetch` API (AJAX)"
    ])

    add_card(slide6, 6.8, 1.6, 5.6, 5.2, "Implementation Workflow & Testing", [
        "1. Startup Init: `init_db()` creates `feedback` table automatically if missing.",
        "2. Public Submission: User fills form -> JS validates -> AJAX POST `/submit-feedback` -> SQL INSERT -> Success Toast.",
        "3. Admin Security: Protected `/admin-dashboard` checks `@login_required` session decorator.",
        "4. Visual Analytics: Admin UI queries `/api/feedback` -> Chart.js renders Rating Distribution & Trend graphs.",
        "5. Verification: Automated integration test suite (`test_app.py`) verified 9/9 test cases successfully passed!"
    ])

    # =========================================================================
    # SLIDE 7: CONCLUSION & SUMMARY
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_background(slide7)
    add_header(slide7, "Conclusion", "Summary & Key Takeaways")

    add_card(slide7, 1.0, 1.6, 11.333, 5.2, "Internship Conclusion & Project Impact", [
        "Successful Internship Execution: Successfully completed the 2-Month Python Developer Internship at The Skybrisk under project guide Govind Bhokare.",
        "Production-Ready Application: Engineered a fully functional, sleek, and robust feedback collection and visual analytics system.",
        "Comprehensive Testing: 100% automated integration test pass rate (`test_app.py`) covering authentication, validation, REST API, and CSV export.",
        "Academic & Professional Fulfillment: Fulfilled all partial requirements for B.Tech (CSE) at Sir Padampat Singhania University (SPSU), Udaipur.",
        "Thank You! Questions & Discussion Welcomed."
    ])

    # Save Presentation
    target_pptx_project = os.path.join(r"C:\Users\bunny\.gemini\antigravity\scratch\OnlineFeedbackCollector", "Online_Feedback_Collector_Presentation_Munukuntla_Bunny_Narayana_Naidu.pptx")
    target_pptx_desktop = r"C:\Users\bunny\OneDrive\Pictures\Desktop\Online_Feedback_Collector_Presentation_Munukuntla_Bunny_Narayana_Naidu.pptx"

    prs.save(target_pptx_project)
    print(f"[OK] Presentation saved to project directory: {target_pptx_project}")

    try:
        prs.save(target_pptx_desktop)
        print(f"[OK] Presentation saved to desktop directory: {target_pptx_desktop}")
    except Exception as e:
        print(f"[NOTE] Desktop save warning: {e}")

if __name__ == "__main__":
    build_presentation()
